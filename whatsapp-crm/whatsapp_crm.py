import csv
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import webbrowser
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

APP_NAME = "WhatsApp CRM"
APP_VERSION = "2.0.0"

# Store data beside the EXE. If running from source, store beside the .py file.
if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

DB = os.path.join(APP_DIR, "whatsapp_crm.db")
CSV_FILE = os.path.join(APP_DIR, "customers.csv")
XLSX_FILE = os.path.join(APP_DIR, "customers.xlsx")
SETTINGS_FILE = os.path.join(APP_DIR, "crm_settings.json")

FIELDS = [
    "name", "phone", "email", "address",
    "pet_name", "pet_type", "purchases", "price", "notes", "next_contact"
]

DEFAULT_LABELS = {
    "name": "Name *",
    "phone": "Phone (+country code)",
    "email": "Email",
    "address": "Address",
    "pet_name": "Pet name",
    "pet_type": "Pet type",
    "purchases": "What customer bought",
    "price": "Price",
    "notes": "Notes",
    "next_contact": "Next contact (YYYY-MM-DD)"
}

DEFAULT_CSV_HEADERS = {
    "name": "Name",
    "phone": "Phone",
    "email": "Email",
    "address": "Address",
    "pet_name": "Pet name",
    "pet_type": "Pet type",
    "purchases": "Purchases",
    "price": "Price",
    "notes": "Notes",
    "next_contact": "Next contact"
}

def load_settings():
    data = {"app_title": APP_NAME, "labels": DEFAULT_LABELS.copy()}
    try:
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            saved = json.load(f)
        if isinstance(saved, dict):
            data["app_title"] = str(saved.get("app_title", APP_NAME))
            labels = saved.get("labels", {})
            if isinstance(labels, dict):
                for k in FIELDS:
                    if labels.get(k):
                        data["labels"][k] = str(labels[k])
    except Exception:
        pass
    return data

settings = load_settings()

def save_settings():
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)

conn = sqlite3.connect(DB)
conn.execute("""CREATE TABLE IF NOT EXISTS customers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT, phone TEXT, email TEXT, address TEXT,
    pet_name TEXT, pet_type TEXT, purchases TEXT, price TEXT,
    notes TEXT, next_contact TEXT
)""")
conn.commit()

def rows():
    return conn.execute(
        "SELECT id,name,phone,email,address,pet_name,pet_type,purchases,price,notes,next_contact "
        "FROM customers ORDER BY name COLLATE NOCASE"
    ).fetchall()

def csv_headers():
    return [DEFAULT_CSV_HEADERS[k] for k in FIELDS]

def sync_csv():
    data = rows()
    with open(CSV_FILE, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(csv_headers())
        for r in data:
            w.writerow(r[1:])

def export_xlsx():
    if Workbook is None:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(csv_headers())
    for r in rows():
        ws.append(list(r[1:]))
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        width = max([12] + [len(str(c.value or "")) + 2 for c in column_cells])
        ws.column_dimensions[letter].width = min(45, width)
    wb.save(XLSX_FILE)
    return True

def save_exports():
    sync_csv()
    export_xlsx()

def normalized_phone(value):
    return re.sub(r"[^0-9]", "", value or "")

def open_whatsapp_for(phone):
    number = normalized_phone(phone)
    if not number:
        messagebox.showwarning("No phone", "Enter the customer phone number including country code.")
        return
    webbrowser.open("https://wa.me/" + number)

def refresh():
    tree.delete(*tree.get_children())
    q = search.get().lower().strip()
    count = 0
    for r in rows():
        vals = r[1:]
        if q and not any(q in str(v or "").lower() for v in vals):
            continue
        tree.insert("", "end", iid=str(r[0]), values=vals)
        count += 1
    status.set(f"{count} customer(s)")

def clear_form():
    selected_id.set("")
    for v in variables.values():
        v.set("")
    name_entry.focus()

def get_selected():
    selected = tree.selection()
    return int(selected[0]) if selected else None

def load_selected(_event=None):
    cid = get_selected()
    if not cid:
        return
    r = conn.execute(
        "SELECT id,name,phone,email,address,pet_name,pet_type,purchases,price,notes,next_contact "
        "FROM customers WHERE id=?", (cid,)
    ).fetchone()
    if not r:
        return
    selected_id.set(str(cid))
    for key, value in zip(FIELDS, r[1:]):
        variables[key].set(value or "")

def validate_date(value):
    if not value:
        return True
    try:
        dt.date.fromisoformat(value)
        return True
    except ValueError:
        return False

def save_customer():
    values = [variables[k].get().strip() for k in FIELDS]
    if not values[0]:
        messagebox.showwarning("Missing name", "Please enter the customer name.")
        return
    if values[-1] and not validate_date(values[-1]):
        messagebox.showwarning("Invalid date", "Use YYYY-MM-DD for the next-contact date.")
        return

    cid = selected_id.get().strip()
    if cid:
        conn.execute(
            "UPDATE customers SET " + ",".join(f"{k}=?" for k in FIELDS) + " WHERE id=?",
            values + [cid]
        )
    else:
        conn.execute(
            "INSERT INTO customers (" + ",".join(FIELDS) + ") VALUES (" +
            ",".join("?" for _ in FIELDS) + ")", values
        )
    conn.commit()
    save_exports()
    refresh()
    clear_form()

def delete_customer():
    cid = get_selected()
    if not cid:
        return
    if messagebox.askyesno("Delete customer", "Delete this customer permanently?"):
        conn.execute("DELETE FROM customers WHERE id=?", (cid,))
        conn.commit()
        save_exports()
        refresh()
        clear_form()

def import_csv():
    path = filedialog.askopenfilename(
        title="Import customer CSV",
        filetypes=[("CSV files", "*.csv"), ("All files", "*.*")]
    )
    if not path:
        return
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,	")
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ";"
            reader = csv.DictReader(f, dialect=dialect)
            if not reader.fieldnames:
                raise ValueError("The CSV has no header row.")

            normalized = {str(h).strip().lower(): h for h in reader.fieldnames}
            aliases = {}
            for k in FIELDS:
                choices = {
                    DEFAULT_CSV_HEADERS[k].lower(),
                    settings["labels"][k].replace(" *", "").lower(),
                    k.lower()
                }
                found = next((normalized[c] for c in choices if c in normalized), None)
                aliases[k] = found

            imported = 0
            for record in reader:
                values = []
                for k in FIELDS:
                    source = aliases.get(k)
                    values.append((record.get(source, "") if source else "") or "")
                if values[0].strip():
                    conn.execute(
                        "INSERT INTO customers (" + ",".join(FIELDS) + ") VALUES (" +
                        ",".join("?" for _ in FIELDS) + ")", values
                    )
                    imported += 1
        conn.commit()
        save_exports()
        refresh()
        messagebox.showinfo("Imported", f"{imported} customer(s) imported successfully.")
    except Exception as exc:
        messagebox.showerror("Import error", str(exc))

def export_files():
    try:
        save_exports()
        messagebox.showinfo(
            "Exported",
            f"Saved in the CRM folder:\n\n{CSV_FILE}\n{XLSX_FILE}"
        )
    except Exception as exc:
        messagebox.showerror("Export error", str(exc))

def notification():
    today = dt.date.today().isoformat()
    due = []
    for r in rows():
        if r[-1] and r[-1] <= today:
            due.append((r[1], r[-1]))
    if due:
        lines = [f"{name} — {date}" for name, date in due]
        messagebox.showinfo(
            "Notifications",
            "Follow-up due today or overdue:\n\n" + "\n".join(lines)
        )
    else:
        messagebox.showinfo("Notifications", "No follow-ups are due today.")

def edit_field_settings():
    dialog = tk.Toplevel(root)
    dialog.title("CRM field settings")
    dialog.geometry("650x600")
    dialog.transient(root)
    dialog.grab_set()

    ttk.Label(
        dialog,
        text="Change the field names to match your business.",
        font=("Segoe UI", 13, "bold")
    ).pack(anchor="w", padx=18, pady=(18, 4))
    ttk.Label(
        dialog,
        text="Example: change “Pet name” to “Company”, “Pet type” to “Industry”, etc.",
        wraplength=590
    ).pack(anchor="w", padx=18, pady=(0, 15))

    title_var = tk.StringVar(value=settings["app_title"])
    ttk.Label(dialog, text="Application title").pack(anchor="w", padx=18)
    ttk.Entry(dialog, textvariable=title_var, width=58).pack(fill="x", padx=18, pady=(4, 12))

    vars_local = {}
    frame = ttk.Frame(dialog)
    frame.pack(fill="both", expand=True, padx=18)

    for i, key in enumerate(FIELDS):
        vars_local[key] = tk.StringVar(value=settings["labels"][key])
        ttk.Label(frame, text=key).grid(row=i, column=0, sticky="w", pady=5)
        ttk.Entry(frame, textvariable=vars_local[key], width=45).grid(row=i, column=1, sticky="ew", padx=10, pady=5)
    frame.columnconfigure(1, weight=1)

    buttons = ttk.Frame(dialog)
    buttons.pack(fill="x", padx=18, pady=18)

    def reset_defaults():
        title_var.set(APP_NAME)
        for key in FIELDS:
            vars_local[key].set(DEFAULT_LABELS[key])

    def apply():
        settings["app_title"] = title_var.get().strip() or APP_NAME
        for key in FIELDS:
            settings["labels"][key] = vars_local[key].get().strip() or DEFAULT_LABELS[key]
        save_settings()
        apply_labels()
        dialog.destroy()

    ttk.Button(buttons, text="Reset defaults", command=reset_defaults).pack(side="left")
    ttk.Button(buttons, text="Cancel", command=dialog.destroy).pack(side="right", padx=(6, 0))
    ttk.Button(buttons, text="Save settings", command=apply).pack(side="right")

def apply_labels():
    root.title(settings["app_title"])
    header_title.config(text=settings["app_title"])
    for key in FIELDS:
        form_labels[key].config(text=settings["labels"][key])

def about():
    messagebox.showinfo(
        "About",
        f"{settings['app_title']}\nVersion {APP_VERSION}\n\n"
        "Customer register, WhatsApp links, CSV/Excel and reminders."
    )

def self_test():
    # Used by the Windows build workflow; never opens a GUI.
    test_conn = sqlite3.connect(":memory:")
    test_conn.execute(
        "CREATE TABLE customers (id INTEGER PRIMARY KEY, name TEXT, phone TEXT)"
    )
    test_conn.execute("INSERT INTO customers(name,phone) VALUES (?,?)", ("Test", "5511999999999"))
    assert test_conn.execute("SELECT COUNT(*) FROM customers").fetchone()[0] == 1
    test_conn.close()
    print("SELF_TEST_OK")

if "--self-test" in sys.argv:
    self_test()
    raise SystemExit(0)

root = tk.Tk()
root.title(settings["app_title"])
root.geometry("1280x760")
root.minsize(1050, 650)

style = ttk.Style()
try:
    style.theme_use("vista")
except Exception:
    pass
style.configure("Treeview", rowheight=28)
style.configure("TButton", padding=7)

selected_id = tk.StringVar()
search = tk.StringVar()
status = tk.StringVar()
variables = {k: tk.StringVar() for k in FIELDS}

header = ttk.Frame(root, padding=12)
header.pack(fill="x")
header_title = ttk.Label(header, text=settings["app_title"], font=("Segoe UI", 22, "bold"))
header_title.pack(side="left")
ttk.Label(
    header,
    text="  Customer register • WhatsApp • purchases • reminders",
    font=("Segoe UI", 10)
).pack(side="left", pady=8)

bar = ttk.Frame(root, padding=(12, 0, 12, 8))
bar.pack(fill="x")
ttk.Label(bar, text="Search").pack(side="left")
ttk.Entry(bar, textvariable=search, width=34).pack(side="left", padx=7)
search.trace_add("write", lambda *_: refresh())
ttk.Button(bar, text="New customer", command=clear_form).pack(side="left", padx=4)
ttk.Button(bar, text="Import CSV", command=import_csv).pack(side="left", padx=4)
ttk.Button(bar, text="Export CSV + Excel", command=export_files).pack(side="left", padx=4)
ttk.Button(bar, text="Notifications", command=notification).pack(side="right", padx=4)
ttk.Button(bar, text="Field settings", command=edit_field_settings).pack(side="right", padx=4)
ttk.Button(bar, text="About", command=about).pack(side="right", padx=4)

main = ttk.Panedwindow(root, orient="horizontal")
main.pack(fill="both", expand=True, padx=12, pady=5)

left = ttk.Frame(main, padding=5)
right = ttk.Frame(main, padding=12)
main.add(left, weight=3)
main.add(right, weight=2)

cols = DEFAULT_CSV_HEADERS
tree = ttk.Treeview(left, columns=FIELDS, show="headings")
for key in FIELDS:
    tree.heading(key, text=DEFAULT_CSV_HEADERS[key])
    tree.column(key, width=115, anchor="w")
tree.column("name", width=155)
tree.column("phone", width=125)
tree.column("purchases", width=155)
tree.column("notes", width=190)
tree.column("next_contact", width=125)

ys = ttk.Scrollbar(left, orient="vertical", command=tree.yview)
xs = ttk.Scrollbar(left, orient="horizontal", command=tree.xview)
tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
tree.grid(row=0, column=0, sticky="nsew")
ys.grid(row=0, column=1, sticky="ns")
xs.grid(row=1, column=0, sticky="ew")
left.rowconfigure(0, weight=1)
left.columnconfigure(0, weight=1)
tree.bind("<<TreeviewSelect>>", load_selected)
tree.bind("<Double-1>", load_selected)

ttk.Label(
    right, text="Customer details", font=("Segoe UI", 15, "bold")
).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))

form_labels = {}
name_entry = None
for i, key in enumerate(FIELDS, 1):
    form_labels[key] = ttk.Label(right, text=settings["labels"][key])
    form_labels[key].grid(row=i, column=0, sticky="w", pady=4)
    ent = ttk.Entry(right, textvariable=variables[key], width=38)
    ent.grid(row=i, column=1, sticky="ew", pady=4)
    if key == "name":
        name_entry = ent

right.columnconfigure(1, weight=1)

buttons = ttk.Frame(right)
buttons.grid(row=12, column=0, columnspan=2, sticky="ew", pady=15)
ttk.Button(buttons, text="Save customer", command=save_customer).pack(side="left", padx=(0, 6))
ttk.Button(
    buttons, text="Open WhatsApp",
    command=lambda: open_whatsapp_for(variables["phone"].get())
).pack(side="left", padx=6)
ttk.Button(buttons, text="Delete", command=delete_customer).pack(side="left", padx=6)

ttk.Label(
    right,
    text="All customer data stays locally beside the app. CSV is updated after every change. "
         "Excel export is available at any time.",
    wraplength=430
).grid(row=13, column=0, columnspan=2, sticky="w", pady=10)

ttk.Label(root, textvariable=status, anchor="w", padding=8).pack(fill="x")

root.bind("<Control-n>", lambda _e: clear_form())
root.bind("<Control-s>", lambda _e: save_customer())
root.bind("<Delete>", lambda _e: delete_customer())

save_exports()
refresh()
name_entry.focus()
root.mainloop()
