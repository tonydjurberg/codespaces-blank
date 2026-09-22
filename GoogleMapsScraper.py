import re, time, threading
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import WebDriverException, StaleElementReferenceException

APP_NAME = "Google Maps Scraper"
DEFAULT_DELAY = 2.5

FIELDS = [
    "Name", "Category", "Address", "Phone", "Website",
    "Google Maps URL", "Rating", "Reviews", "Search Query", "Scraped At"
]

def clean(value):
    return re.sub(r"\s+", " ", value or "").strip()

def first_match(text, patterns):
    for p in patterns:
        m = re.search(p, text, re.I)
        if m:
            return clean(m.group(1))
    return ""

def parse_card(card, query):
    text = clean(card.text)
    if not text:
        return None

    name = ""
    try:
        h = card.find_element(By.CSS_SELECTOR, "div.fontHeadlineSmall")
        name = clean(h.text)
    except Exception:
        pass
    if not name:
        try:
            a = card.find_element(By.CSS_SELECTOR, 'a[href*="/maps/place/"]')
            name = clean(a.get_attribute("aria-label") or a.text)
        except Exception:
            pass
    if not name:
        lines = [clean(x) for x in card.text.splitlines() if clean(x)]
        name = lines[0] if lines else ""

    maps_url = ""
    try:
        a = card.find_element(By.CSS_SELECTOR, 'a[href*="/maps/place/"]')
        maps_url = a.get_attribute("href") or ""
    except Exception:
        pass

    website = ""
    try:
        a = card.find_element(By.CSS_SELECTOR, 'a[data-item-id="authority"]')
        website = a.get_attribute("href") or ""
    except Exception:
        pass

    phone = ""
    for el in card.find_elements(By.CSS_SELECTOR, 'button, a'):
        try:
            label = el.get_attribute("aria-label") or ""
            if re.search(r"(phone|telefon|call|ring)", label, re.I):
                phone = re.sub(r"^(Phone|Telefon|Call|Ring):?\s*", "", label, flags=re.I)
                break
        except Exception:
            pass
    if not phone:
        phone = first_match(text, [
            r"(\+\d[\d\s().-]{7,}\d)",
            r"(0\d[\d\s().-]{7,}\d)"
        ])

    rating = first_match(text, [r"([0-5](?:[.,]\d)?)\s*(?:stars?|stjärnor?)"])
    reviews = first_match(text, [r"\((\d[\d\s.,]*)\)"])
    category = ""
    address = ""

    lines = [clean(x) for x in card.text.splitlines() if clean(x)]
    for line in lines[1:]:
        if line == name or re.fullmatch(r"[0-5](?:[.,]\d)?", line):
            continue
        if not category and len(line) < 80 and not re.search(r"\d{2,}", line):
            category = line
            continue
        if not address and len(line) < 180:
            address = line
            break

    return {
        "Name": name,
        "Category": category,
        "Address": address,
        "Phone": phone,
        "Website": website,
        "Google Maps URL": maps_url,
        "Rating": rating,
        "Reviews": reviews,
        "Search Query": query,
        "Scraped At": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }

class App:
    def __init__(self, root):
        self.root = root
        root.title(APP_NAME)
        root.geometry("760x540")
        root.minsize(680, 480)

        frm = ttk.Frame(root, padding=18)
        frm.pack(fill="both", expand=True)

        ttk.Label(frm, text=APP_NAME, font=("Segoe UI", 20, "bold")).pack(anchor="w")
        ttk.Label(frm, text="Slow browser-based search • no API key", font=("Segoe UI", 10)).pack(anchor="w", pady=(0,18))

        row = ttk.Frame(frm); row.pack(fill="x", pady=5)
        ttk.Label(row, text="Search:", width=16).pack(side="left")
        self.query = tk.StringVar(value="real estate agents Malmö Sweden")
        ttk.Entry(row, textvariable=self.query).pack(side="left", fill="x", expand=True)

        row = ttk.Frame(frm); row.pack(fill="x", pady=5)
        ttk.Label(row, text="Maximum results:", width=16).pack(side="left")
        self.max_results = tk.IntVar(value=50)
        ttk.Spinbox(row, from_=1, to=500, textvariable=self.max_results, width=8).pack(side="left")
        ttk.Label(row, text="Delay per scroll (seconds):").pack(side="left", padx=(30,8))
        self.delay = tk.DoubleVar(value=DEFAULT_DELAY)
        ttk.Spinbox(row, from_=1.0, to=15.0, increment=0.5, textvariable=self.delay, width=8).pack(side="left")

        row = ttk.Frame(frm); row.pack(fill="x", pady=5)
        ttk.Label(row, text="Output:", width=16).pack(side="left")
        self.output = tk.StringVar(value=str(Path.home() / "Desktop" / "GoogleMaps_Results.xlsx"))
        ttk.Entry(row, textvariable=self.output).pack(side="left", fill="x", expand=True)
        ttk.Button(row, text="Browse", command=self.browse).pack(side="left", padx=(8,0))

        self.start_btn = ttk.Button(frm, text="START SCRAPE", command=self.start)
        self.start_btn.pack(fill="x", pady=(18,8), ipady=7)

        self.status = tk.StringVar(value="Ready.")
        ttk.Label(frm, textvariable=self.status).pack(anchor="w")

        self.log = tk.Text(frm, height=14, wrap="word")
        self.log.pack(fill="both", expand=True, pady=(10,0))
        self.driver = None

    def browse(self):
        p = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel workbook","*.xlsx")],
            initialfile="GoogleMaps_Results.xlsx"
        )
        if p:
            self.output.set(p)

    def write_log(self, msg):
        self.root.after(0, lambda: (self.log.insert("end", msg + "\n"), self.log.see("end")))

    def set_status(self, msg):
        self.root.after(0, lambda: self.status.set(msg))

    def start(self):
        q = clean(self.query.get())
        if not q:
            messagebox.showwarning(APP_NAME, "Enter a Google Maps search.")
            return
        self.start_btn.config(state="disabled")
        self.log.delete("1.0","end")
        threading.Thread(target=self.scrape, args=(q,), daemon=True).start()

    def scrape(self, query):
        results = {}
        try:
            self.set_status("Starting Chrome...")
            opts = Options()
            opts.add_argument("--start-maximized")
            opts.add_argument("--disable-notifications")
            opts.add_argument("--lang=en-US")
            self.driver = webdriver.Chrome(options=opts)

            url = "https://www.google.com/maps/search/" + quote_plus(query)
            self.write_log("Opening: " + url)
            self.driver.get(url)
            time.sleep(4)

            feed = None
            for _ in range(20):
                try:
                    feed = self.driver.find_element(By.CSS_SELECTOR, 'div[role="feed"]')
                    break
                except Exception:
                    time.sleep(0.5)
            if not feed:
                raise RuntimeError("Google Maps results feed was not found. Check the browser window.")

            delay = max(1.0, float(self.delay.get()))
            target = max(1, int(self.max_results.get()))
            stable = 0
            last_count = 0

            while len(results) < target and stable < 8:
                cards = feed.find_elements(By.CSS_SELECTOR, ':scope > div')
                for card in cards:
                    try:
                        item = parse_card(card, query)
                        if item and item["Name"]:
                            key = item["Google Maps URL"] or (item["Name"] + "|" + item["Address"])
                            results[key] = item
                    except StaleElementReferenceException:
                        continue
                    if len(results) >= target:
                        break

                self.set_status(f"Collected {len(results)} / {target}")
                self.write_log(f"Collected {len(results)} result(s). Scrolling...")
                self.driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight", feed)
                time.sleep(delay)

                if len(results) == last_count:
                    stable += 1
                else:
                    stable = 0
                    last_count = len(results)

            data = list(results.values())[:target]
            self.save_xlsx(data, self.output.get())
            self.set_status(f"Done — {len(data)} result(s)")
            self.write_log(f"Saved {len(data)} result(s) to: {self.output.get()}")
            self.root.after(0, lambda: messagebox.showinfo(APP_NAME, f"Done.\n\n{len(data)} results saved."))
        except Exception as e:
            self.write_log("ERROR: " + repr(e))
            self.set_status("Error.")
            self.root.after(0, lambda: messagebox.showerror(APP_NAME, str(e)))
        finally:
            try:
                if self.driver:
                    self.driver.quit()
            except Exception:
                pass
            self.driver = None
            self.root.after(0, lambda: self.start_btn.config(state="normal"))

    def save_xlsx(self, data, path):
        out = Path(path).expanduser()
        out.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Google Maps Results"
        ws.append(FIELDS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9EAF7")
        for item in data:
            ws.append([item.get(f, "") for f in FIELDS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = [34,24,50,22,45,70,10,12,35,20]
        for i,w in enumerate(widths,1):
            ws.column_dimensions[chr(64+i)].width = w
        wb.save(out)

if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
